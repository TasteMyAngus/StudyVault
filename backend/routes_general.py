import json
import sqlite3
import traceback
import uuid
from pathlib import Path

from fastapi import APIRouter, File, Form, HTTPException, UploadFile

from .config import DB_PATH, UPLOADS_DIR
from .state import pipeline

router = APIRouter()

@router.get("/health")
async def health_check():
    return {"status": "ok", "service": "StudyVault API"}


@router.get("/")
async def root():
    return {
        "app": "StudyVault API",
        "version": "0.1.0",
        "status": "running",
        "roadmap": "W0-W2: Ingestion + DB schema"
    }


@router.get("/courses")
async def get_courses():
    conn = sqlite3.connect(str(DB_PATH))
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()
    
    cursor.execute('SELECT course_id, name, term, created_at FROM courses ORDER BY created_at DESC')
    rows = cursor.fetchall()
    conn.close()
    
    return [dict(row) for row in rows]


@router.post("/courses")
async def create_course(name: str = Form(...), term: str = Form(None)):
    course_id = str(uuid.uuid4())
    conn = sqlite3.connect(str(DB_PATH))
    cursor = conn.cursor()
    
    cursor.execute('''
        INSERT INTO courses (course_id, name, term)
        VALUES (?, ?, ?)
    ''', (course_id, name, term))
    
    conn.commit()
    conn.close()
    
    print(f"Created course: {name} ({course_id}) in {DB_PATH}")
    
    return {"course_id": course_id, "name": name, "term": term}


@router.get("/courses/{course_id}")
async def get_course(course_id: str):
    conn = sqlite3.connect(str(DB_PATH))
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()
    
    cursor.execute('SELECT * FROM courses WHERE course_id = ?', (course_id,))
    row = cursor.fetchone()
    conn.close()
    
    if not row:
        raise HTTPException(status_code=404, detail="Course not found")
    
    return dict(row)


@router.delete("/courses/{course_id}")
async def delete_course(course_id: str):
    conn = sqlite3.connect(str(DB_PATH))
    cursor = conn.cursor()
    
    cursor.execute('SELECT course_id FROM courses WHERE course_id = ?', (course_id,))
    if not cursor.fetchone():
        conn.close()
        raise HTTPException(status_code=404, detail="Course not found")
    
    cursor.execute('SELECT doc_id FROM documents WHERE course_id = ?', (course_id,))
    doc_ids = [row[0] for row in cursor.fetchall()]
    
    for doc_id in doc_ids:
        cursor.execute('SELECT version_id FROM document_versions WHERE doc_id = ?', (doc_id,))
        version_ids = [row[0] for row in cursor.fetchall()]
        
        for version_id in version_ids:
            cursor.execute('SELECT chunk_id FROM chunks WHERE version_id = ?', (version_id,))
            chunk_ids = [row[0] for row in cursor.fetchall()]
            
            for chunk_id in chunk_ids:
                cursor.execute('DELETE FROM embeddings WHERE chunk_id = ?', (chunk_id,))
            
            cursor.execute('DELETE FROM chunks WHERE version_id = ?', (version_id,))
        
        cursor.execute('DELETE FROM document_versions WHERE doc_id = ?', (doc_id,))
    
    cursor.execute('DELETE FROM documents WHERE course_id = ?', (course_id,))
    
    cursor.execute('SELECT conversation_id FROM conversations WHERE course_id = ?', (course_id,))
    conv_ids = [row[0] for row in cursor.fetchall()]
    for conv_id in conv_ids:
        cursor.execute('DELETE FROM messages WHERE conversation_id = ?', (conv_id,))
        cursor.execute('DELETE FROM retrieval_traces WHERE message_id IN (SELECT message_id FROM messages WHERE conversation_id = ?)', (conv_id,))
    cursor.execute('DELETE FROM conversations WHERE course_id = ?', (course_id,))
    
    cursor.execute('DELETE FROM courses WHERE course_id = ?', (course_id,))
    
    conn.commit()
    conn.close()
    
    print(f"Deleted course and all related data: {course_id}")
    return {"message": "Course deleted successfully", "course_id": course_id}


@router.get("/courses/{course_id}/documents")
async def get_documents(course_id: str):
    conn = sqlite3.connect(str(DB_PATH))
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()
    
    cursor.execute('''
        SELECT doc_id, title, doc_type, category, created_at 
        FROM documents 
        WHERE course_id = ? 
        ORDER BY category, created_at DESC
    ''', (course_id,))
    
    rows = cursor.fetchall()
    conn.close()
    
    return [dict(row) for row in rows]


@router.get("/courses/{course_id}/topics")
async def get_topics(course_id: str):
    conn = sqlite3.connect(str(DB_PATH))
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    cursor.execute('SELECT course_id FROM courses WHERE course_id = ?', (course_id,))
    if not cursor.fetchone():
        conn.close()
        raise HTTPException(status_code=404, detail="Course not found")

    cursor.execute('''
        SELECT t.topic_id, t.label, t.created_at,
               COUNT(tc.chunk_id) AS chunk_count
        FROM topics t
        LEFT JOIN topic_chunks tc ON t.topic_id = tc.topic_id
        WHERE t.course_id = ?
        GROUP BY t.topic_id
        ORDER BY t.created_at DESC
    ''', (course_id,))

    rows = cursor.fetchall()
    conn.close()

    return [dict(row) for row in rows]


@router.post("/courses/{course_id}/topics")
async def create_topic(course_id: str, label: str = Form(...), chunk_ids_json: str = Form(None)):
    label = label.strip()
    if not label:
        raise HTTPException(status_code=400, detail="Topic label cannot be empty")

    conn = sqlite3.connect(str(DB_PATH))
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    cursor.execute('SELECT course_id FROM courses WHERE course_id = ?', (course_id,))
    if not cursor.fetchone():
        conn.close()
        raise HTTPException(status_code=404, detail="Course not found")

    topic_id = str(uuid.uuid4())
    cursor.execute('''
        INSERT INTO topics (topic_id, course_id, label)
        VALUES (?, ?, ?)
    ''', (topic_id, course_id, label))

    cursor.execute('''
        INSERT OR IGNORE INTO topic_mastery (course_id, topic_id, mastery_score)
        VALUES (?, ?, 0.5)
    ''', (course_id, topic_id))

    mapped_chunks = 0
    if chunk_ids_json:
        try:
            chunk_ids = json.loads(chunk_ids_json)
        except json.JSONDecodeError:
            conn.close()
            raise HTTPException(status_code=400, detail="chunk_ids_json must be valid JSON")

        if not isinstance(chunk_ids, list):
            conn.close()
            raise HTTPException(status_code=400, detail="chunk_ids_json must be a JSON list")

        chunk_ids = [c for c in chunk_ids if isinstance(c, str) and c.strip()]
        if chunk_ids:
            placeholders = ",".join(["?"] * len(chunk_ids))
            cursor.execute(f'''
                SELECT c.chunk_id
                FROM chunks c
                JOIN document_versions dv ON c.version_id = dv.version_id
                JOIN documents d ON dv.doc_id = d.doc_id
                WHERE d.course_id = ? AND c.chunk_id IN ({placeholders})
            ''', [course_id] + chunk_ids)
            valid_chunk_ids = [row[0] for row in cursor.fetchall()]

            for chunk_id in valid_chunk_ids:
                cursor.execute('''
                    INSERT OR IGNORE INTO topic_chunks (topic_id, chunk_id, weight)
                    VALUES (?, ?, 1.0)
                ''', (topic_id, chunk_id))
            mapped_chunks = len(valid_chunk_ids)

    conn.commit()
    conn.close()

    return {
        "topic_id": topic_id,
        "label": label,
        "mapped_chunks": mapped_chunks
    }


@router.post("/courses/{course_id}/topics/{topic_id}/attach-documents")
async def attach_documents_to_topic(course_id: str, topic_id: str, doc_ids_json: str = Form(...)):
    conn = sqlite3.connect(str(DB_PATH))
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    cursor.execute('SELECT course_id FROM courses WHERE course_id = ?', (course_id,))
    if not cursor.fetchone():
        conn.close()
        raise HTTPException(status_code=404, detail="Course not found")

    cursor.execute('SELECT topic_id FROM topics WHERE topic_id = ? AND course_id = ?', (topic_id, course_id))
    if not cursor.fetchone():
        conn.close()
        raise HTTPException(status_code=404, detail="Topic not found for this course")

    try:
        parsed_doc_ids = json.loads(doc_ids_json)
    except json.JSONDecodeError:
        conn.close()
        raise HTTPException(status_code=400, detail="doc_ids_json must be valid JSON")

    if not isinstance(parsed_doc_ids, list):
        conn.close()
        raise HTTPException(status_code=400, detail="doc_ids_json must be a JSON list")

    selected_doc_ids = [str(doc_id).strip() for doc_id in parsed_doc_ids if str(doc_id).strip()]
    if not selected_doc_ids:
        conn.close()
        raise HTTPException(status_code=400, detail="Please select at least one document")

    placeholders = ",".join(["?"] * len(selected_doc_ids))
    cursor.execute(f'''
        SELECT doc_id
        FROM documents
        WHERE course_id = ? AND doc_id IN ({placeholders})
    ''', [course_id] + selected_doc_ids)
    valid_doc_ids = [row["doc_id"] for row in cursor.fetchall()]
    if not valid_doc_ids:
        conn.close()
        raise HTTPException(status_code=400, detail="No valid course documents selected")

    doc_placeholders = ",".join(["?"] * len(valid_doc_ids))
    cursor.execute(f'''
        SELECT c.chunk_id
        FROM chunks c
        JOIN document_versions dv ON c.version_id = dv.version_id
        JOIN (
            SELECT doc_id, MAX(uploaded_at) AS latest_uploaded_at
            FROM document_versions
            WHERE doc_id IN ({doc_placeholders})
            GROUP BY doc_id
        ) latest ON latest.doc_id = dv.doc_id AND latest.latest_uploaded_at = dv.uploaded_at
        ORDER BY dv.doc_id, c.chunk_index ASC
    ''', valid_doc_ids)
    valid_chunk_ids = [row["chunk_id"] for row in cursor.fetchall()]

    if not valid_chunk_ids:
        conn.close()
        raise HTTPException(status_code=400, detail="No chunks found for selected documents")

    chunk_placeholders = ",".join(["?"] * len(valid_chunk_ids))
    cursor.execute(f'''
        SELECT chunk_id
        FROM topic_chunks
        WHERE topic_id = ? AND chunk_id IN ({chunk_placeholders})
    ''', [topic_id] + valid_chunk_ids)
    existing_chunk_ids = {row["chunk_id"] for row in cursor.fetchall()}

    new_links = 0
    for chunk_id in valid_chunk_ids:
        cursor.execute('''
            INSERT OR IGNORE INTO topic_chunks (topic_id, chunk_id, weight)
            VALUES (?, ?, 1.0)
        ''', (topic_id, chunk_id))
        if chunk_id not in existing_chunk_ids:
            new_links += 1

    conn.commit()
    conn.close()

    return {
        "topic_id": topic_id,
        "selected_doc_count": len(valid_doc_ids),
        "mapped_chunks": len(valid_chunk_ids),
        "new_links": new_links
    }


@router.post("/courses/{course_id}/topics/auto")
async def auto_detect_topics(course_id: str):
    # Build topics from section titles when they exist.
    conn = sqlite3.connect(str(DB_PATH))
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    cursor.execute('SELECT course_id FROM courses WHERE course_id = ?', (course_id,))
    if not cursor.fetchone():
        conn.close()
        raise HTTPException(status_code=404, detail="Course not found")

    cursor.execute('''
        SELECT topic_id, label
        FROM topics
        WHERE course_id = ?
    ''', (course_id,))
    existing_topics = cursor.fetchall()
    label_to_topic_id = {row['label'].strip().lower(): row['topic_id'] for row in existing_topics}

    cursor.execute('''
        SELECT c.chunk_id, c.section_title
        FROM chunks c
        JOIN document_versions dv ON c.version_id = dv.version_id
        JOIN documents d ON dv.doc_id = d.doc_id
        WHERE d.course_id = ?
          AND c.section_title IS NOT NULL
          AND TRIM(c.section_title) != ''
    ''', (course_id,))

    rows = cursor.fetchall()

    created_topics = 0
    mapped_chunks = 0

    for row in rows:
        chunk_id = row['chunk_id']
        raw_label = row['section_title'].strip()
        label_key = raw_label.lower()

        topic_id = label_to_topic_id.get(label_key)
        if not topic_id:
            topic_id = str(uuid.uuid4())
            cursor.execute('''
                INSERT INTO topics (topic_id, course_id, label)
                VALUES (?, ?, ?)
            ''', (topic_id, course_id, raw_label))
            cursor.execute('''
                INSERT OR IGNORE INTO topic_mastery (course_id, topic_id, mastery_score)
                VALUES (?, ?, 0.5)
            ''', (course_id, topic_id))
            label_to_topic_id[label_key] = topic_id
            created_topics += 1

        cursor.execute('''
            INSERT OR IGNORE INTO topic_chunks (topic_id, chunk_id, weight)
            VALUES (?, ?, 1.0)
        ''', (topic_id, chunk_id))
        mapped_chunks += 1

    conn.commit()
    conn.close()

    return {
        "created_topics": created_topics,
        "mapped_chunks": mapped_chunks
    }


@router.post("/courses/{course_id}/documents")
async def upload_document(course_id: str, title: str = Form(...), category: str = Form("General"), file: UploadFile = File(...)):
    doc_id = str(uuid.uuid4())
    
    file_ext = Path(file.filename).suffix.lower().strip('.')
    if file_ext not in ["pdf", "txt", "md", "docx", "pptx"]:
        raise HTTPException(status_code=400, detail=f"Unsupported file type: {file_ext}")
    
    file_path = UPLOADS_DIR / f"{doc_id}_{file.filename}"
    with open(file_path, 'wb') as f:
        content = await file.read()
        f.write(content)
    
    conn = sqlite3.connect(str(DB_PATH))
    cursor = conn.cursor()
    cursor.execute('SELECT course_id FROM courses WHERE course_id = ?', (course_id,))
    if not cursor.fetchone():
        conn.close()
        raise HTTPException(status_code=404, detail="Course not found")
    
    cursor.execute('''
        INSERT INTO documents (doc_id, course_id, title, doc_type, category)
        VALUES (?, ?, ?, ?, ?)
    ''', (doc_id, course_id, title, file_ext, category if category else 'General'))
    
    conn.commit()
    conn.close()
    
    try:
        print(f"Starting ingestion for {file.filename} (doc_id={doc_id}, type={file_ext})")
        result = pipeline.ingest(str(file_path), doc_id, course_id, file_ext, title)
        print(f"Ingestion successful: {result}")
        return {
            "doc_id": doc_id,
            "title": title,
            "file_path": str(file_path),
            "ingestion_result": result
        }
    except Exception as e:
        import traceback
        print(f"Ingestion error: {str(e)}")
        print(traceback.format_exc())
        raise HTTPException(status_code=500, detail=f"Ingestion failed: {str(e)}")


@router.get("/documents/{doc_id}/content")
async def get_document_content(doc_id: str):
    conn = sqlite3.connect(str(DB_PATH))
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()
    
    cursor.execute('SELECT title, doc_type FROM documents WHERE doc_id = ?', (doc_id,))
    doc_row = cursor.fetchone()
    if not doc_row:
        conn.close()
        raise HTTPException(status_code=404, detail="Document not found")
    
    doc_type = doc_row['doc_type']
    title = doc_row['title']
    
    cursor.execute('SELECT file_path FROM document_versions WHERE doc_id = ? ORDER BY uploaded_at DESC LIMIT 1', (doc_id,))
    version_row = cursor.fetchone()
    conn.close()
    
    if not version_row:
        raise HTTPException(status_code=404, detail="Document file not found")
    
    file_path = version_row['file_path']
    
    try:
        if doc_type == 'docx':
            from docx import Document
            from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
            from docx.oxml.text.paragraph import CT_P
            from docx.oxml.table import CT_Tbl
            from docx.table import _Cell, Table
            from docx.text.paragraph import Paragraph
            
            doc = Document(file_path)
            
            # Keep paragraphs and tables in the order they appear in the file.
            html_content = []
            current_list = None  # Track if we're in a list
            list_items = []
            
            for block in doc.element.body:
                if isinstance(block, CT_P):
                    para = Paragraph(block, doc)
                    if not para.text.strip():
                        continue
                    
                    # Treat Word list styles as real HTML lists.
                    is_list_item = False
                    list_type = None
                    if para.style and para.style.name:
                        style_name = para.style.name.lower()
                        if 'list bullet' in style_name or para.text.strip().startswith('•'):
                            is_list_item = True
                            list_type = 'ul'
                        elif 'list number' in style_name or (len(para.text) > 0 and para.text[0].isdigit() and '. ' in para.text[:5]):
                            is_list_item = True
                            list_type = 'ol'
                    
                    if is_list_item:
                        if current_list != list_type:
                            if current_list and list_items:
                                html_content.append(f'<{current_list}>')
                                html_content.extend(list_items)
                                html_content.append(f'</{current_list}>')
                                list_items = []
                            current_list = list_type
                        
                        
                        text = para.text.strip()
                        
                        if text.startswith('•'):
                            text = text[1:].strip()
                        elif text[0].isdigit() and '. ' in text[:5]:
                            text = text.split('. ', 1)[1] if '. ' in text else text
                        
                        li_html = '<li>'
                        for run in para.runs:
                            run_text = run.text
                            if run.bold and run.italic:
                                run_text = f'<strong><em>{run_text}</em></strong>'
                            elif run.bold:
                                run_text = f'<strong>{run_text}</strong>'
                            elif run.italic:
                                run_text = f'<em>{run_text}</em>'
                            elif run.underline:
                                run_text = f'<u>{run_text}</u>'
                            li_html += run_text
                        li_html += '</li>'
                        list_items.append(li_html)
                    else:
                        if current_list and list_items:
                            html_content.append(f'<{current_list}>')
                            html_content.extend(list_items)
                            html_content.append(f'</{current_list}>')
                            list_items = []
                            current_list = None
                        
                        if para.style.name.startswith('Heading'):
                            level = para.style.name[-1] if para.style.name[-1].isdigit() else '1'
                            html_content.append(f'<h{level}>{para.text}</h{level}>')
                        else:
                            # Preserve basic inline formatting and tab spacing.
                            para_html = '<p>'
                            for run in para.runs:
                                text = run.text.replace('\t', '&nbsp;&nbsp;&nbsp;&nbsp;')  # Convert tabs to spaces
                                if run.bold and run.italic:
                                    text = f'<strong><em>{text}</em></strong>'
                                elif run.bold:
                                    text = f'<strong>{text}</strong>'
                                elif run.italic:
                                    text = f'<em>{text}</em>'
                                elif run.underline:
                                    text = f'<u>{text}</u>'
                                para_html += text
                            para_html += '</p>'
                            html_content.append(para_html)
                
                elif isinstance(block, CT_Tbl):
                    if current_list and list_items:
                        html_content.append(f'<{current_list}>')
                        html_content.extend(list_items)
                        html_content.append(f'</{current_list}>')
                        list_items = []
                        current_list = None
                    
                    table = Table(block, doc)
                    html_content.append('<table>')
                    for i, row in enumerate(table.rows):
                        html_content.append('<tr>')
                        for cell in row.cells:
                            cell_text = cell.text.strip()
                            tag = 'th' if i == 0 else 'td'
                            html_content.append(f'<{tag}>{cell_text}</{tag}>')
                        html_content.append('</tr>')
                    html_content.append('</table>')
            
            if current_list and list_items:
                html_content.append(f'<{current_list}>')
                html_content.extend(list_items)
                html_content.append(f'</{current_list}>')
            
            content = '\n'.join(html_content)
            is_html = True
        elif doc_type in ['txt', 'md']:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
            is_html = False
        elif doc_type == 'pdf':
            from pypdf import PdfReader
            reader = PdfReader(file_path)
            content = '\n\n'.join([page.extract_text() for page in reader.pages])
            is_html = False
        else:
            raise HTTPException(status_code=400, detail="Unsupported document type")
        
        return {
            "doc_id": doc_id,
            "title": title,
            "doc_type": doc_type,
            "content": content,
            "is_html": is_html
        }
    except Exception as e:
        print(f"Error reading document: {str(e)}")
        raise HTTPException(status_code=500, detail=f"Failed to read document: {str(e)}")


